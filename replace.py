from docx import Document
import os
from logger import create_log

def replace_text_in_docx(file_path, replacements):
    doc = Document(file_path)
    log_entries = []

    # Remplacement dans les paragraphes
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for find, replace, col in replacements:
                if not col and find in run.text:
                    log_entries.append(f"Paragraphe : '{find}' -> '{replace}'")
                    run.text = run.text.replace(find, replace)

    # Remplacement dans les tableaux
    for table in doc.tables:
        header = table.rows[0]
        column_map = {cell.text.strip(): idx for idx, cell in enumerate(header.cells)}

        for find, replace, col in replacements:
            if col and col in column_map:
                col_idx = column_map[col]
                for row in table.rows[1:]:
                    if col_idx < len(row.cells):
                        cell = row.cells[col_idx]
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if find in run.text:
                                    log_entries.append(f"Table[{col}] : '{find}' -> '{replace}'")
                                    run.text = run.text.replace(find, replace)
            elif not col:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                if find in run.text:
                                    log_entries.append(f"Table[All] : '{find}' -> '{replace}'")
                                    run.text = run.text.replace(find, replace)

    # Sauvegarde du nouveau fichier
    new_path = file_path.replace(".docx", "_modified.docx")
    doc.save(new_path)

    # Création du log
    create_log(file_path, replacements)
    return new_path
